"""Criação de Word, Excel, PowerPoint, PDF, TXT e Markdown por comando natural.

O pedido do usuário determina o destino; o modelo só fornece o CONTEÚDO.
Sem sobrescrita, execução de scripts ou gravações silenciosas fora do perfil.
"""
import json
from pathlib import Path
import re
import unicodedata
from urllib.parse import unquote
from document_writer import (FileRequest, create_file as write_basic, generate_contents,
                             validate_title, known_folder, _reserve_unique_file)

FORMAT = {
    "word": "word", "docx": "word", "documento": "word",
    "excel": "excel", "xlsx": "excel", "planilha": "excel",
    "powerpoint": "powerpoint", "pptx": "powerpoint", "apresentacao": "powerpoint",
    "slides": "powerpoint", "pdf": "pdf",
    "texto": "texto", "txt": "texto", "arquivo": "texto",
    "markdown": "markdown", "md": "markdown",
}
EXT = {"word": ".docx", "excel": ".xlsx", "powerpoint": ".pptx",
       "pdf": ".pdf", "texto": ".txt", "markdown": ".md"}
VERBS = r"(?:crie|cria|criar|gere|gera|gerar|faca|fazer|monte|montar|produza)"
DEST = (r"(?:na|no|em|para a|para o)\s+(?:pasta\s+)?"
        r"(Documentos|Documents|Downloads|Área de Trabalho|Area de Trabalho|Desktop)")
PLAN_SCHEMA = {
    "type": "object", "additionalProperties": False,
    "properties": {
        "titulo": {"type": "string"},
        "itens": {"type": "array", "items": {"type": "string"}, "maxItems": 12},
    }, "required": ["titulo", "itens"]
}


def folded(text):
    chars = unicodedata.normalize("NFKD", text.casefold())
    return "".join(c for c in chars if not unicodedata.combining(c))


def file_intent(message):
    """Só cria após pedido operacional, não quando a pessoa menciona criar algo."""
    clean = folded(message.strip())
    clean = re.sub(r"^yuna[, ]+", "", clean)
    if clean.startswith("/criar "):
        return True
    if re.match(r"^(?:como|por que|porque|e se|se eu|o que e|nao|nunca)\b", clean):
        return False
    request = re.match(
        r"^(?:(?:por favor|ei yuna)[, ]+)?"
        r"(?:(?:voce|vc)\s+)?"
        r"(?:(?:pode|consegue|poderia)\s+)?"
        r"(?:crie|cria|faca|monte|gere|produza)\b",
        clean
    ) or re.match(
        r"^(?:quero|preciso)\s+que\s+(?:voce|a yuna)?\s*"
        r"(?:crie|cria|faca|monte|gere)\b", clean
    )
    return bool(request and re.search(
        r"\b(?:word|docx|documento|planilha|excel|powerpoint|pptx|"
        r"apresentacao|slides|pdf|arquivo|texto|txt|markdown)\b", clean
    ))


def parse_request(message, home=None):
    """Cria plano somente de texto literal, nunca usando localização sugerida pelo LLM."""
    text = message.strip()
    if not file_intent(text):
        return None
    normalized = folded(text)
    verb = re.search(r"\b" + VERBS + r"\b", normalized)
    after = normalized[verb.end():] if verb else normalized
    fmt = next((FORMAT[key] for key in FORMAT if re.search(r"\b" + re.escape(key) + r"\b", after)), "texto")
    # Suporte a /criar excel|pptx|pdf e pastas conhecidas.
    quoted = re.search(r'["“]([^"”]{1,100})["”]', text)
    title = quoted.group(1) if quoted else None
    explicit = re.search(
        r"\b(?:chamado|chamada|intitulado|intitulada|nome(?:\s+de)?)\s+"
        r"(.{2,150})", text, re.I
    )
    destination = re.search(DEST + r"[\s.!?]*$", text, re.I)
    dest_name = destination.group(1) if destination else "Documentos"

    if title is None and explicit:
        title = explicit.group(1)
        if destination:
            title = title[:max(0, title.lower().rfind(destination.group(0).lower()))]
        title = title.strip(" '\".,!?")
    # /criar excel Gastos mensais | Documentos | Assunto
    desc = ""
    if normalized.startswith("/criar "):
        slash = re.match(r"/criar\s+([^\s]+)\s+(.+)", text, re.I)
        if slash:
            fmt = FORMAT.get(folded(slash.group(1)), fmt)
            parts = [p.strip() for p in slash.group(2).split("|")]
            if parts and parts[0]:
                title = title or parts[0]
                if len(parts) > 1 and parts[1]:
                    dest_name = parts[1]
                if len(parts) > 2:
                    desc = parts[2][:500]
    if not title:
        title = {
            "word": "Documento da Yuna", "excel": "Planilha da Yuna",
            "powerpoint": "Apresentação da Yuna", "pdf": "Relatório da Yuna",
            "texto": "Anotações da Yuna", "markdown": "Notas da Yuna"
        }[fmt]
    directory = known_folder(dest_name, home)
    # Se o usuário mencionar um caminho absoluto desconhecido, não fingir que salvou lá.
    if not destination and not normalized.startswith("/criar ") and re.search(
            r"[a-z]:\\|\\\\|/home/|na pasta \S+", text, re.I):
        raise ValueError("Para um caminho específico, use /criar FORMATO NOME | Documentos, "
                         "ou escolha a pasta explicitamente no seletor 📄.")
    if not desc:
        desc = text[:500]
    return FileRequest(fmt, validate_title(title), directory, desc)


def _list_from_model(request, ask):
    """JSON é conteúdo validado, não comandos. Só um conjunto pequeno de células/slides."""
    messages = [
        {"role": "system", "content": (
            "Você redige o CONTEÚDO de um arquivo solicitado em português brasileiro. "
            "Retorne JSON exatamente com titulo:string e itens:array de 3 a 8 strings. "
            "Para Excel, itens são linhas descritivas ou dados ilustrativos, sempre "
            "identificados como exemplos; não invente valores reais do usuário. "
            "Para PowerPoint, cada item é o título e a descrição de um slide. "
            "Não produza comandos, URLs, macros, fórmulas arbitrárias ou scripts."
        )},
        {"role": "user", "content": f"Arquivo {request.kind}; título: {request.title}; tema: {request.description}"}
    ]
    raw = ask(messages, schema=PLAN_SCHEMA)
    try:
        data = json.loads(raw)
    except (TypeError, ValueError) as exc:
        raise ValueError("O Ollama não retornou conteúdo estruturado válido.") from exc
    if (not isinstance(data, dict) or set(data) != {"titulo", "itens"}
            or not isinstance(data["titulo"], str) or len(data["titulo"]) > 150
            or not isinstance(data["itens"], list) or not 1 <= len(data["itens"]) <= 12
            or not all(isinstance(x, str) and 3 <= len(x) <= 450 for x in data["itens"])):
        raise ValueError("O modelo retornou uma estrutura de documento inválida.")
    return data


def create_requested(request, ask):
    if request.kind in ("word", "texto", "markdown"):
        content = generate_contents(request, ask)
        return write_basic(request, content)
    path = None
    if request.kind in ("excel", "powerpoint"):
        data = _list_from_model(request, ask)
    else:
        data = None
    if request.kind == "excel":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError as exc:
            raise RuntimeError("Instale a biblioteca gratuita: py -3 -m pip install openpyxl") from exc
        book = Workbook()
        sheet = book.active
        sheet.title = "Resumo"
        sheet.append([request.title, "Descrição"])
        for index, item in enumerate(data["itens"], 1):
            # Strings não confiáveis começam com apóstrofo para impedir fórmulas.
            value = item if not item.startswith(("=", "+", "-", "@")) else "'" + item
            sheet.append([f"Item {index}", value])
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 75
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="443366")
            cell.alignment = Alignment(wrap_text=True)
        path = _reserve_unique_file(request.directory, request.title, EXT[request.kind])
        try:
            book.save(path)
            return path
        except Exception:
            path.unlink(missing_ok=True)
            raise
    if request.kind == "powerpoint":
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("Instale a biblioteca gratuita: py -3 -m pip install python-pptx") from exc
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[0])
        slide.shapes.title.text = request.title
        slide.placeholders[1].text = "Criado localmente pela Yuna · Conteúdo gerado pelo Ollama"
        for entry in data["itens"][:8]:
            slide = pres.slides.add_slide(pres.slide_layouts[1])
            head, sep, body = entry.partition(":")
            slide.shapes.title.text = head[:90] if sep else f"Parte {len(pres.slides) - 1}"
            slide.placeholders[1].text = body.strip() if sep else entry
        path = _reserve_unique_file(request.directory, request.title, EXT[request.kind])
        try:
            pres.save(path)
            return path
        except Exception:
            path.unlink(missing_ok=True)
            raise
    if request.kind == "pdf":
        content = generate_contents(request, ask)
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.enums import TA_LEFT
        except ImportError as exc:
            raise RuntimeError("Instale a biblioteca gratuita: py -3 -m pip install reportlab") from exc
        styles = getSampleStyleSheet()
        font_path = Path("C:/Windows/Fonts/arial.ttf")
        if font_path.is_file():
            pdfmetrics.registerFont(TTFont("YunaArial", str(font_path)))
            styles["Normal"].fontName = "YunaArial"
            styles["Title"].fontName = "YunaArial"
        from xml.sax.saxutils import escape
        story = [Paragraph(escape(request.title), styles["Title"]), Spacer(1, 12)]
        for line in content.splitlines():
            if line.strip():
                story.append(Paragraph(escape(line.strip()), styles["Normal"]))
                story.append(Spacer(1, 8))
        path = _reserve_unique_file(request.directory, request.title, EXT[request.kind])
        try:
            SimpleDocTemplate(str(path), pagesize=A4).build(story)
            return path
        except Exception:
            path.unlink(missing_ok=True)
            raise
    raise ValueError("Formato não implementado.")
