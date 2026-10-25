"""Arquivos reais: parser sem modelo, destinos explícitos e não sobrescrita."""
import tempfile
from pathlib import Path
import unittest
from unittest.mock import Mock, patch
import os

from document_writer import (
    FileRequest, create_file, generate_contents, known_folder,
    looks_like_file_request, parse_file_request, validate_title
)


class DocumentWriterTests(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.TemporaryDirectory()
        self.home = Path(self.root.name)
        for name in ("Documents", "Desktop", "Downloads"):
            (self.home / name).mkdir()

    def tearDown(self):
        self.root.cleanup()

    def test_parse_natural_word_and_requested_folder(self):
        request = parse_file_request(
            "Yuna, crie um documento Word chamado 'Diário da Yuna' na pasta Documentos",
            home=self.home
        )
        self.assertEqual(request.kind, "word")
        self.assertEqual(request.title, "Diário da Yuna")
        self.assertEqual(request.directory, self.home / "Documents")
        req = parse_file_request(
            "Cria um Word chamado Contas 2026 na Área de Trabalho", home=self.home
        )
        self.assertEqual(req.title, "Contas 2026")
        self.assertEqual(req.directory, self.home / "Desktop")
        self.assertTrue(looks_like_file_request("Crie um documento do Word"))

    def test_parse_explicit_slash(self):
        request = parse_file_request(
            "/criar word Diário da Yuna | Downloads | Primeira página de diário",
            home=self.home
        )
        self.assertEqual(request.directory, self.home / "Downloads")
        self.assertEqual(request.description, "Primeira página de diário")
        request = parse_file_request("/criar texto Lista | Documentos", home=self.home)
        self.assertEqual(request.kind, "texto")

    def test_no_hallucinated_document_for_incomplete_request(self):
        with self.assertRaises(ValueError):
            parse_file_request("Crie um documento Word para mim", home=self.home)
        self.assertIsNone(parse_file_request("Como faço um Word?", home=self.home))
        with self.assertRaises(ValueError):
            known_folder("C:\\Windows", self.home)

    def test_windows_known_folder_redirected_to_onedrive(self):
        if os.name != "nt":
            self.skipTest("Windows known folders require Windows")
        redirected = self.home / "OneDrive" / "Desktop"
        redirected.mkdir(parents=True)
        with patch("document_writer._windows_known_folder", return_value=redirected):
            self.assertEqual(known_folder("Área de Trabalho"), redirected)
        with patch("document_writer._windows_known_folder", return_value=redirected):
            self.assertEqual(known_folder("Desktop"), redirected)

    def test_unique_name_prevents_overwrite(self):
        request = FileRequest("texto", "Diário", self.home / "Documents")
        first = create_file(request, "Uma primeira entrada.")
        second = create_file(request, "Uma segunda entrada.")
        self.assertEqual(first.name, "Diário.txt")
        self.assertEqual(second.name, "Diário (2).txt")
        self.assertEqual(first.read_text(encoding="utf-8"), "Uma primeira entrada.\n")
        self.assertEqual(second.read_text(encoding="utf-8"), "Uma segunda entrada.\n")

    def test_document_folder_is_explicit_not_from_model(self):
        request = FileRequest("word", "Relatório", self.home / "Desktop",
                              "Descreva o projeto.")
        model = Mock(return_value="Abertura do projeto Yuna.\n\nSegunda seção de teste.")
        self.assertIn("Abertura", generate_contents(request, model))
        messages = model.call_args.args[0]
        self.assertNotIn(str(self.home), repr(messages))
        self.assertIn("não comandos do sistema", messages[0]["content"])

    def test_protect_invalid_titles(self):
        for name in ("", "../..", "CON", "aux.txt"):
            with self.assertRaises(ValueError):
                validate_title(name)
        self.assertNotIn("/", validate_title("Nome / perigoso"))

    def test_does_not_create_in_nonexistent_folder(self):
        target = FileRequest("word", "Teste", self.home / "Inexistente")
        with self.assertRaisesRegex(ValueError, "não existe"):
            create_file(target, "Conteúdo")

    def test_real_docx_contents_and_no_overwrite(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed in this environment")
        request = FileRequest("word", "Diário da Yuna", self.home / "Documents")
        first = create_file(request, "Primeira página.\n\nMais uma anotação.")
        second = create_file(request, "Segunda página.")
        self.assertEqual(first.suffix, ".docx")
        self.assertEqual(second.name, "Diário da Yuna (2).docx")
        first_doc = Document(first)
        self.assertEqual(first_doc.paragraphs[0].text, "Diário da Yuna")
        self.assertIn("Primeira página.", [p.text for p in first_doc.paragraphs])
        self.assertEqual(Document(second).paragraphs[-1].text, "Segunda página.")


if __name__ == "__main__":
    unittest.main()
