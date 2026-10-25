"""Yuna 0.12: criação real de arquivos office e parser natural, tudo offline."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock

from office_agent import file_intent, parse_request, create_requested, folded


class OfficeAgentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.home = Path(self.tmp.name)
        for name in ("Documents", "Downloads", "Desktop"):
            (self.home / name).mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    @staticmethod
    def fake_model(messages, schema=None):
        if schema is not None:
            return json.dumps({"titulo": "Projeto da Yuna", "itens": [
                "Introdução: metas do projeto", "Etapa 1: primeira melhoria",
                "Etapa 2: validação técnica"
            ]})
        return "Este documento resume os objetivos do projeto Yuna.\n\nAqui está a primeira etapa."

    def test_natural_requests_with_default_title_and_destination(self):
        request = parse_request("Crie uma apresentação PowerPoint sobre redes", self.home)
        self.assertEqual(request.kind, "powerpoint")
        self.assertEqual(request.title, "Apresentação da Yuna")
        self.assertEqual(request.directory, self.home / "Documents")
        request = parse_request("Monte uma planilha Excel sobre despesas na Área de Trabalho", self.home)
        self.assertEqual(request.kind, "excel")
        self.assertEqual(request.directory, self.home / "Desktop")
        request = parse_request("Crie um arquivo", self.home)
        self.assertEqual(request.kind, "texto")

    def test_explicit_names_and_no_inference_from_questions(self):
        req = parse_request(
            "Crie um documento Word chamado 'Diário da Yuna' em Documentos",
            self.home
        )
        self.assertEqual(req.title, "Diário da Yuna")
        self.assertEqual(req.kind, "word")
        self.assertFalse(file_intent("Como criar um PowerPoint?"))
        self.assertFalse(file_intent("Não crie um arquivo"))
        self.assertTrue(file_intent("Faça um PDF para mim"))

    def test_creates_all_supported_office_formats(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        for format_name, command in [
            ("word", "Crie um Word"),
            ("excel", "Crie uma planilha Excel"),
            ("powerpoint", "Crie um PowerPoint"),
            ("pdf", "Crie um PDF"),
            ("texto", "Crie um arquivo"),
            ("markdown", "/criar markdown Notas da Yuna | Documentos | ideias"),
        ]:
            with self.subTest(format_name):
                req = parse_request(command, self.home)
                path = create_requested(req, self.fake_model)
                self.assertTrue(path.exists())
                self.assertGreater(path.stat().st_size, 0)
                if format_name == "word":
                    self.assertTrue(Document(path).paragraphs)
                if format_name == "excel":
                    self.assertEqual(load_workbook(path).active.max_row, 4)
                if format_name == "powerpoint":
                    self.assertEqual(len(Presentation(path).slides), 4)
                if format_name == "pdf":
                    self.assertEqual(path.read_bytes()[:4], b"%PDF")

    def test_does_not_overwrite_existing_document(self):
        req = parse_request("Crie um PowerPoint", self.home)
        first = create_requested(req, self.fake_model)
        second = create_requested(req, self.fake_model)
        self.assertEqual(first.name, "Apresentação da Yuna.pptx")
        self.assertEqual(second.name, "Apresentação da Yuna (2).pptx")
        self.assertNotEqual(first, second)

    def test_llm_cannot_choose_output_directory(self):
        req = parse_request("Crie uma planilha Excel", self.home)
        def model(messages, schema=None):
            self.assertNotIn(str(self.home), repr(messages))
            return json.dumps({"titulo": "Tenta mudar pasta", "itens": [
                "=HYPERLINK(\"malicioso\")", "Ideia de teste"
            ]})
        path = create_requested(req, model)
        self.assertEqual(path.parent, self.home / "Documents")
        from openpyxl import load_workbook
        self.assertNotEqual(load_workbook(path).active["B2"].data_type, "f")

    def test_rejects_unknown_explicit_directory(self):
        with self.assertRaises(ValueError):
            parse_request("Crie um Word na pasta C:\\Windows\\System32", self.home)


if __name__ == "__main__":
    unittest.main()
