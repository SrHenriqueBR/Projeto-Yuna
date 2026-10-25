"""Yuna 0.13: integration tests for bounded multi-step tasks, no network."""
import json
from pathlib import Path
import tempfile
import unittest

from agent_brain import (
    TaskPlan, plan_request, describe_plan, execute_plan, verify_file,
)
from office_agent import create_requested


SOURCE = ("Inteligência artificial: área da ciência da computação.\n"
          "Fonte: https://pt.wikipedia.org/wiki/Intelig%C3%AAncia_artificial\n"
          "Algoritmos: técnicas computacionais.\n"
          "Fonte: https://pt.wikipedia.org/wiki/Algoritmo")


class BrainTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.home = Path(self.tmp.name)
        for folder in ("Documents", "Downloads", "Desktop"):
            (self.home / folder).mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    @staticmethod
    def ask(messages, schema=None):
        assert any("Fonte:" in msg["content"] for msg in messages), "O contexto da fonte precisa ser passado ao Ollama"
        if schema:
            return json.dumps({
                "titulo": "Inteligência artificial",
                "itens": [
                    "O que é: estudo dos sistemas computacionais",
                    "Exemplos: ferramentas de processamento de texto",
                    "Limites: uso responsável das tecnologias",
                ]
            })
        return ("Este documento apresenta a inteligência artificial como campo de "
                "pesquisa computacional.\n\nEstes resultados são resumidos e "
                "não substituem fontes técnicas detalhadas.")

    def test_plans_mentioned_goal_and_destination_without_model(self):
        request = ("Pesquise sobre inteligência artificial, crie um PowerPoint "
                   "com os resultados e salve nos Documentos")
        plan = plan_request(request, self.home)
        self.assertIsInstance(plan, TaskPlan)
        self.assertEqual(plan.topic, "inteligencia artificial")
        self.assertEqual(plan.output.kind, "powerpoint")
        self.assertEqual(plan.output.directory, self.home / "Documents")
        self.assertIn("não executado", describe_plan(plan))

    def test_area_de_trabalho_and_other_formats(self):
        request = plan_request(
            "Yuna, pesquise sobre astronomia e depois crie um documento PDF na Área de Trabalho",
            self.home)
        self.assertEqual(request.output.kind, "pdf")
        self.assertEqual(request.output.directory, self.home / "Desktop")
        self.assertIsNone(plan_request("Crie um Word sobre astronomia", self.home))

    def test_rejects_hypothetical_or_personal_research(self):
        self.assertIsNone(plan_request(
            "Como pesquisar sobre astronomia e criar um PowerPoint?", self.home))
        self.assertIsNone(plan_request(
            "Eu vou pesquisar astronomia e criar um PowerPoint depois", self.home))
        with self.assertRaisesRegex(ValueError, "dados pessoais"):
            plan_request(
                "Pesquise sobre minha conta, crie um PDF em Documentos", self.home)

    def test_writes_powerpoint_and_sources_and_progress(self):
        plan = plan_request(
            "Pesquise sobre inteligência artificial, crie um PowerPoint e salve nos Documentos",
            self.home)
        progress = []
        response = execute_plan(
            plan, self.ask, progress=progress.append,
            lookup=lambda query: SOURCE, create=create_requested)
        self.assertEqual(len(progress), 4)
        self.assertTrue(progress[0].startswith("1/4"))
        self.assertTrue(progress[-1].startswith("4/4"))
        self.assertIn("✅ Plano concluído", response)
        files = list((self.home / "Documents").glob("*.pptx"))
        self.assertEqual(len(files), 1)
        from pptx import Presentation
        slides = Presentation(files[0]).slides
        self.assertEqual(len(slides), 5)  # capa + 3 slides + fontes
        self.assertEqual(slides[-1].shapes.title.text, "Fontes consultadas")
        self.assertIn("Fonte", response)

    def test_word_and_excel_source_attachments(self):
        from docx import Document
        from openpyxl import load_workbook
        for kind, command in [
            ("word", "Pesquise sobre astronomia, crie um Word em Documentos"),
            ("excel", "Pesquise sobre astronomia, crie uma planilha Excel em Documentos"),
        ]:
            with self.subTest(kind):
                plan = plan_request(command, self.home)
                outcome = execute_plan(
                    plan, self.ask, lookup=lambda query: SOURCE)
                self.assertIn("Concluído", outcome.replace("concluído", "Concluído"))
                file = next((self.home / "Documents").glob("Pesquisa*." + {
                    "word": "docx", "excel": "xlsx"}[kind]))
                verify_file(file, plan.output)
                if kind == "word":
                    self.assertIn("Fontes consultadas",
                                  [p.text for p in Document(file).paragraphs])
                else:
                    self.assertIn("Fontes", load_workbook(file).sheetnames)

    def test_no_document_on_source_failure(self):
        plan = plan_request(
            "Pesquise sobre astronomia, crie PowerPoint em Documentos",
            self.home)
        with self.assertRaisesRegex(RuntimeError, "não retornou fontes"):
            execute_plan(plan, self.ask, lookup=lambda query: "Nenhum resultado.")
        self.assertEqual(list((self.home / "Documents").iterdir()), [])

    def test_dry_run_does_not_call_any_services(self):
        plan = plan_request(
            "Pesquise sobre astronomia, crie PDF nos Documentos", self.home)
        self.assertIn("pesquisa", describe_plan(plan).lower())
        self.assertEqual(list((self.home / "Documents").iterdir()), [])


if __name__ == "__main__":
    unittest.main()
